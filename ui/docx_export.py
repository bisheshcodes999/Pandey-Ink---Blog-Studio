"""Converts the backend's Markdown into a Word document (.docx).

Just an export option for the frontend - the backend still writes
Markdown to disk as the source of truth (that's what Saved Blogs,
per-section regenerate/splice, and citation rendering all work
against). This module just gives a second, ready-to-open format for
the same content when "Download Word (.docx)" gets clicked in the
Article tab - converts the exact same Markdown that's already on
screen.

Went with a small hand-written Markdown -> docx walk instead of a
generic converter on purpose: the backend only ever emits a known,
fairly narrow set of Markdown constructs (see WORKER_SYSTEM in
backend/nodes/worker.py and generate_and_place_images() in
backend/reducer/images.py), so a purpose-built parser covers it fine
without pulling in something heavier (pandoc) that'd need a separate
install.
"""
from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path
from typing import List

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

_ACCENT_HEX = "2563EB"
_MUTED_RGB = RGBColor(0x64, 0x74, 0x8B)

_IMAGE_LINE_RE = re.compile(r"^!\[(?P<alt>[^\]]*)\]\((?P<src>[^)]+)\)\s*$")
_FAILED_IMAGE_RE = re.compile(r"^> \*\*\[IMAGE GENERATION FAILED\]\*\* ?(?P<caption>.*)$")
_HR_RE = re.compile(r"^(-{3,}|\*{3,}|_{3,})$")
_BULLET_RE = re.compile(r"^(\s*)[-*+]\s+(.*)$")
_NUMBERED_RE = re.compile(r"^(\s*)\d+[.)]\s+(.*)$")
_TABLE_SEP_RE = re.compile(r"^\|?\s*:?-+:?\s*(\|\s*:?-+:?\s*)*\|?$")

# Inline spans: bold, italic, code, links - longest-rule-first so
# "**bold**" doesn't get chewed up by the italic rule first.
_INLINE_RE = re.compile(r"(\*\*.+?\*\*|`.+?`|\[.+?\]\(.+?\)|\*.+?\*)")
_BOLD_RE = re.compile(r"^\*\*(.+)\*\*$")
_ITALIC_RE = re.compile(r"^\*(.+)\*$")
_CODE_RE = re.compile(r"^`(.+)`$")
_LINK_RE = re.compile(r"^\[(.+?)\]\((.+?)\)$")


def _add_hyperlink(paragraph, url: str, text: str) -> None:
    """Adds an actual clickable hyperlink run. python-docx doesn't have
    a hyperlink API built in, so this hand-builds the bit of OOXML Word
    expects (a relationship to the URL, plus a <w:hyperlink> run
    pointing at it)."""
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    run_el = OxmlElement("w:r")
    run_props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), _ACCENT_HEX)
    run_props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_props.append(underline)
    run_el.append(run_props)

    text_el = OxmlElement("w:t")
    text_el.text = text
    run_el.append(text_el)
    hyperlink.append(run_el)
    paragraph._p.append(hyperlink)


def _add_inline_runs(paragraph, text: str) -> None:
    """Pulls **bold**, *italic*, `code`, and [text](url) spans out of
    one line and adds each as its own formatted run."""
    pos = 0
    for match in _INLINE_RE.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos:match.start()])
        token = match.group(0)

        bold_match = _BOLD_RE.match(token)
        code_match = _CODE_RE.match(token)
        link_match = _LINK_RE.match(token)
        italic_match = _ITALIC_RE.match(token)

        if bold_match:
            paragraph.add_run(bold_match.group(1)).bold = True
        elif code_match:
            run = paragraph.add_run(code_match.group(1))
            run.font.name = "Consolas"
        elif link_match:
            link_text, url = link_match.groups()
            _add_hyperlink(paragraph, url, link_text)
        elif italic_match:
            paragraph.add_run(italic_match.group(1)).italic = True
        else:
            paragraph.add_run(token)
        pos = match.end()

    if pos < len(text):
        paragraph.add_run(text[pos:])


def _add_horizontal_rule(doc: Document) -> None:
    paragraph = doc.add_paragraph()
    p_properties = paragraph._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:color"), "CCCCCC")
    border.append(bottom)
    p_properties.append(border)


def _bullet_style_for(indent: int) -> str:
    level = min(indent // 2, 2)
    return ["List Bullet", "List Bullet 2", "List Bullet 3"][level]


def _resolve_local_image(src: str) -> Path:
    """Same resolution rule as ui/blog_viewer.py's image inlining, so
    the Word export and the in-app article viewer agree on where an
    image actually lives on disk."""
    return Path(src.strip().lstrip("./")).resolve()


def markdown_to_docx_bytes(md: str) -> bytes:
    """Converts the backend's Markdown (headings, lists, bold/italic/
    code, links, local images + captions, the failed-image callout,
    tables, fenced code blocks) into a formatted .docx and returns its
    bytes."""
    doc = Document()
    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Calibri"
    normal_style.font.size = Pt(11)

    lines = md.replace("\r\n", "\n").split("\n")
    total = len(lines)
    index = 0
    in_code_block = False
    code_buffer: List[str] = []

    def flush_code_block() -> None:
        if not code_buffer:
            return
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.space_after = Pt(6)
        run = paragraph.add_run("\n".join(code_buffer))
        run.font.name = "Consolas"
        run.font.size = Pt(9.5)
        code_buffer.clear()

    while index < total:
        line = lines[index]
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code_block:
                flush_code_block()
            in_code_block = not in_code_block
            index += 1
            continue

        if in_code_block:
            code_buffer.append(line)
            index += 1
            continue

        if not stripped:
            index += 1
            continue

        # Failed-image callout - a fixed 4-field blockquote block from
        # generate_and_place_images() when a Gemini call fails.
        if stripped.startswith("> **[IMAGE GENERATION FAILED]**"):
            caption_match = _FAILED_IMAGE_RE.match(stripped)
            caption = caption_match.group("caption").strip() if caption_match else ""
            index += 1
            while index < total and lines[index].strip().startswith(">"):
                index += 1
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.3)
            run = paragraph.add_run(f"[Image could not be generated: {caption or 'unavailable'}]")
            run.italic = True
            run.font.color.rgb = _MUTED_RGB
            continue

        # Local image ref, inlined as a real picture (falls back to a
        # bracketed note if the file isn't actually on disk).
        image_match = _IMAGE_LINE_RE.match(stripped)
        if image_match:
            alt = image_match.group("alt")
            image_path = _resolve_local_image(image_match.group("src"))
            if image_path.exists():
                try:
                    doc.add_picture(str(image_path), width=Inches(5.5))
                    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                except Exception:
                    fallback = doc.add_paragraph()
                    fallback.add_run(f"[Image: {alt or image_path.name}]").italic = True
            else:
                fallback = doc.add_paragraph()
                fallback.add_run(f"[Image: {alt or image_path.name}]").italic = True

            index += 1
            # A caption line (a lone *italic* line) right after an
            # image - same convention the article viewer's CSS uses.
            if index < total and _ITALIC_RE.match(lines[index].strip()):
                caption_text = _ITALIC_RE.match(lines[index].strip()).group(1)
                caption_paragraph = doc.add_paragraph()
                caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                caption_run = caption_paragraph.add_run(caption_text)
                caption_run.italic = True
                caption_run.font.size = Pt(9.5)
                caption_run.font.color.rgb = _MUTED_RGB
                index += 1
            continue

        if stripped.startswith("#### "):
            doc.add_heading(stripped[5:].strip(), level=4)
            index += 1
            continue
        if stripped.startswith("### "):
            doc.add_heading(stripped[4:].strip(), level=3)
            index += 1
            continue
        if stripped.startswith("## "):
            doc.add_heading(stripped[3:].strip(), level=2)
            index += 1
            continue
        if stripped.startswith("# "):
            doc.add_heading(stripped[2:].strip(), level=1)
            index += 1
            continue

        if _HR_RE.match(stripped):
            _add_horizontal_rule(doc)
            index += 1
            continue

        # Generic blockquote (anything else starting with "> ").
        if stripped.startswith(">"):
            block_lines = [stripped.lstrip(">").strip()]
            index += 1
            while index < total and lines[index].strip().startswith(">"):
                block_lines.append(lines[index].strip().lstrip(">").strip())
                index += 1
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.3)
            _add_inline_runs(paragraph, " ".join(part for part in block_lines if part))
            for run in paragraph.runs:
                run.italic = True
            continue

        bullet_match = _BULLET_RE.match(line)
        if bullet_match:
            indent, text = bullet_match.groups()
            paragraph = doc.add_paragraph(style=_bullet_style_for(len(indent)))
            _add_inline_runs(paragraph, text)
            index += 1
            continue

        numbered_match = _NUMBERED_RE.match(line)
        if numbered_match:
            paragraph = doc.add_paragraph(style="List Number")
            _add_inline_runs(paragraph, numbered_match.group(2))
            index += 1
            continue

        # Simple pipe-table support.
        if stripped.startswith("|") and stripped.endswith("|"):
            table_lines = [stripped]
            index += 1
            while index < total and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            rows = [
                [cell.strip() for cell in row.strip("|").split("|")]
                for row in table_lines
                if not _TABLE_SEP_RE.match(row)
            ]
            if rows:
                table = doc.add_table(rows=len(rows), cols=len(rows[0]))
                table.style = "Light Grid Accent 1"
                for row_index, row in enumerate(rows):
                    for col_index, cell_text in enumerate(row):
                        if col_index < len(table.columns):
                            table.cell(row_index, col_index).text = cell_text
            continue

        paragraph = doc.add_paragraph()
        _add_inline_runs(paragraph, stripped)
        index += 1

    flush_code_block()

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
